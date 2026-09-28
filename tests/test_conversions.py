import pytest

from document_converter import ConversionError, Converter, FORMATS
from document_converter.converter import unique_path

from .conftest import PHRASE, SOURCES, extract_text

converter = Converter()


def _pairs():
    for kind in SOURCES:
        for target in converter.targets_for(FORMATS[kind]):
            yield kind, target.key


@pytest.mark.parametrize("kind,target", list(_pairs()))
def test_every_offered_conversion_works(sample, tmp_path, kind, target):
    src = sample(kind)
    outputs = converter.convert(src, target, tmp_path / "out")
    assert outputs, "no files written"
    for out in outputs:
        assert out.exists() and out.stat().st_size > 0
        assert out.suffix == "." + FORMATS[target].extension
    text = "\n".join(extract_text(o) for o in outputs)
    # A CSV of a document holds its table, not the paragraphs around it.
    expected = "North" if target in ("csv", "tsv") and kind in ("md", "html", "docx", "pptx") else PHRASE
    assert expected in text or expected.replace(" ", " ") in text, f"content lost in {kind} to {target}"


def test_targets_exclude_own_format():
    for fmt in FORMATS.values():
        assert fmt not in converter.targets_for(fmt)


def test_core_formats_always_available():
    for key in ("pdf", "docx", "odt", "rtf", "md", "html", "txt", "epub", "xlsx", "csv", "pptx"):
        assert converter.can_write(FORMATS[key]), key
        assert converter.can_read(FORMATS[key]), key


def test_existing_output_is_not_overwritten(sample, tmp_path):
    src = sample("md")
    first = converter.convert(src, "html", tmp_path)[0]
    second = converter.convert(src, "html", tmp_path)[0]
    assert first != second and second.name == "sample (1).html"
    third = converter.convert(src, "html", tmp_path, overwrite=True)[0]
    assert third == first


def test_markdown_output_uses_md_extension(sample, tmp_path):
    out = converter.convert(sample("docx"), "md", tmp_path)[0]
    assert out.suffix == ".md"
    assert "# Report" in out.read_text(encoding="utf-8")


def test_markdown_to_text_has_no_html_tags(sample, tmp_path):
    out = converter.convert(sample("md"), "txt", tmp_path)[0]
    text = out.read_text(encoding="utf-8")
    assert "<p>" not in text and "<h1>" not in text and PHRASE in text


def test_pdf_output_contains_text(sample, tmp_path):
    out = converter.convert(sample("docx"), "pdf", tmp_path)[0]
    assert PHRASE in extract_text(out)


def test_multi_sheet_workbook_to_csv_writes_one_file_per_sheet(sample, tmp_path):
    outputs = converter.convert(sample("xlsx"), "csv", tmp_path)
    assert sorted(o.name for o in outputs) == ["sample - Detail.csv", "sample - Summary.csv"]


def test_csv_to_xlsx_types_numbers_but_keeps_codes(sample, tmp_path):
    from openpyxl import load_workbook

    out = converter.convert(sample("csv"), "xlsx", tmp_path)[0]
    ws = load_workbook(str(out)).active
    assert ws["B2"].value == 1200.5
    assert ws["B3"].value == "007"


def test_document_tables_become_spreadsheet_rows(sample, tmp_path):
    from openpyxl import load_workbook

    out = converter.convert(sample("docx"), "xlsx", tmp_path)[0]
    rows = list(load_workbook(str(out)).active.iter_rows(values_only=True))
    assert rows[0] == ("Region", "Sales")
    assert rows[1] == ("North", 120)


def test_ods_round_trip(sample, tmp_path):
    ods = converter.convert(sample("xlsx"), "ods", tmp_path)[0]
    back = converter.convert(ods, "xlsx", tmp_path / "back")[0]
    assert PHRASE in extract_text(back)


def test_non_utf8_text_is_decoded(tmp_path):
    src = tmp_path / "legacy.txt"
    src.write_bytes("Caf\xe9 cr\xe8me".encode("cp1252"))
    out = converter.convert(src, "html", tmp_path)[0]
    assert "Café crème" in out.read_text(encoding="utf-8")


def test_scanned_pdf_gives_a_clear_error(tmp_path):
    from pypdf import PdfWriter

    blank = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(612, 792)
    writer.write(str(blank))
    with pytest.raises(ConversionError, match="no selectable text"):
        converter.convert(blank, "docx", tmp_path)


def test_unknown_extension_is_rejected(tmp_path):
    src = tmp_path / "file.xyz"
    src.write_text("hi")
    with pytest.raises(ConversionError, match="Unsupported"):
        converter.convert(src, "pdf", tmp_path)


def test_legacy_formats_explain_libreoffice(tmp_path, monkeypatch):
    from document_converter import engines

    monkeypatch.setenv("DOCCONV_NO_LIBREOFFICE", "1")
    engines.refresh()
    try:
        src = tmp_path / "old.doc"
        src.write_bytes(b"\xd0\xcf\x11\xe0")
        with pytest.raises(ConversionError, match="LibreOffice"):
            converter.convert(src, "pdf", tmp_path)
    finally:
        monkeypatch.delenv("DOCCONV_NO_LIBREOFFICE")
        engines.refresh()


def test_unique_path(tmp_path):
    p = tmp_path / "a.txt"
    assert unique_path(p) == p
    p.write_text("x")
    assert unique_path(p).name == "a (1).txt"


def test_long_tables_flow_across_pdf_pages(tmp_path):
    from openpyxl import Workbook
    from pypdf import PdfReader

    src = tmp_path / "long.xlsx"
    wb = Workbook()
    for r in range(1, 201):
        wb.active.append([r, f"Item {r}"])
    wb.save(str(src))
    out = converter.convert(src, "pdf", tmp_path)[0]
    reader = PdfReader(str(out))
    assert len(reader.pages) > 1
    assert "Item 200" in extract_text(out)
    assert reader.pages[0].extract_text().strip(), "first page should not be blank"


def test_wide_sheets_print_landscape(tmp_path):
    from openpyxl import Workbook
    from pypdf import PdfReader

    src = tmp_path / "wide.xlsx"
    wb = Workbook()
    wb.active.append([f"Col {i}" for i in range(16)])
    wb.save(str(src))
    page = PdfReader(str(converter.convert(src, "pdf", tmp_path)[0])).pages[0]
    assert page.mediabox.width > page.mediabox.height


def test_document_title_survives_markdown(sample, tmp_path):
    from docx import Document

    src = tmp_path / "titled.docx"
    doc = Document()
    doc.add_heading("Annual Report", 0)
    doc.add_paragraph("Body")
    doc.save(str(src))
    out = converter.convert(src, "md", tmp_path)[0]
    assert out.read_text(encoding="utf-8").startswith("# Annual Report")
