import zipfile
from pathlib import Path

import pytest

PHRASE = "Quarterly revenue"


def make_md(path: Path) -> Path:
    path.write_text(f"# Report\n\n{PHRASE} grew **12%** this year.\n\n- North\n- South\n\n"
                    "| Region | Sales |\n|---|---|\n| North | 120 |\n| South | 95 |\n", encoding="utf-8")
    return path


def make_txt(path: Path) -> Path:
    path.write_text(f"Report\n\n{PHRASE} grew 12% this year.\nSecond line.\n\nCafé and naïve text.\n",
                    encoding="utf-8")
    return path


def make_html(path: Path) -> Path:
    path.write_text(f"<html><body><h1>Report</h1><p>{PHRASE} grew 12%.</p>"
                    "<table><tr><th>Region</th><th>Sales</th></tr><tr><td>North</td><td>120</td></tr></table>"
                    "</body></html>", encoding="utf-8")
    return path


def make_docx(path: Path) -> Path:
    from docx import Document

    doc = Document()
    doc.add_heading("Report", 1)
    doc.add_paragraph(f"{PHRASE} grew 12% this year.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Region", "Sales"
    table.cell(1, 0).text, table.cell(1, 1).text = "North", "120"
    doc.save(str(path))
    return path


def make_xlsx(path: Path) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Metric", "Value"])
    ws.append([PHRASE, 1200.5])
    ws.append(["Units", 42])
    wb.create_sheet("Detail").append(["Region", "Sales"])
    wb.save(str(path))
    return path


def make_csv(path: Path) -> Path:
    path.write_text(f"Metric,Value\n{PHRASE},1200.5\nCode,007\n", encoding="utf-8")
    return path


def make_pptx(path: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Report"
    slide.placeholders[1].text_frame.text = f"{PHRASE} grew 12%"
    slide2 = prs.slides.add_slide(prs.slide_layouts[5])
    slide2.shapes.title.text = "Numbers"
    rows = slide2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    rows.cell(0, 0).text, rows.cell(0, 1).text = "Region", "Sales"
    rows.cell(1, 0).text, rows.cell(1, 1).text = "North", "120"
    prs.save(str(path))
    return path


def make_pdf(path: Path) -> Path:
    import typst

    source = path.with_suffix(".typ")
    source.write_text(f"= Report\n\n{PHRASE} grew 12% this year. This paragraph is long enough "
                      "to wrap across more than one line in the generated PDF document.\n", encoding="utf-8")
    typst.compile(str(source), output=str(path))
    source.unlink()
    return path


SOURCES = {
    "md": make_md,
    "txt": make_txt,
    "html": make_html,
    "docx": make_docx,
    "xlsx": make_xlsx,
    "csv": make_csv,
    "pptx": make_pptx,
    "pdf": make_pdf,
}


@pytest.fixture
def sample(tmp_path):
    def build(kind: str) -> Path:
        folder = tmp_path / "in"
        folder.mkdir(exist_ok=True)
        return SOURCES[kind](folder / f"sample.{kind}")
    return build


def extract_text(path: Path) -> str:
    """Best-effort text of any output file, for checking content survived."""
    ext = path.suffix.lower().lstrip(".")
    if ext == "pdf":
        from pypdf import PdfReader
        return "\n".join(p.extract_text() or "" for p in PdfReader(str(path)).pages)
    if ext == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(str(path))
        return "\n".join(str(c) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row if c)
    if ext in ("docx", "pptx", "odt", "ods", "epub"):
        with zipfile.ZipFile(path) as zf:
            return "\n".join(zf.read(n).decode("utf-8", "replace") for n in zf.namelist()
                             if n.endswith((".xml", ".xhtml", ".html")))
    return path.read_bytes().decode("utf-8", "replace")
