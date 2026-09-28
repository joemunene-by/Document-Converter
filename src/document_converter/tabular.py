"""Reading and writing spreadsheets, and moving tables to and from documents."""

from __future__ import annotations

import csv
import datetime as dt
import html
import re
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, List, NamedTuple, Optional

from .textio import read_text


class Sheet(NamedTuple):
    name: str
    rows: List[List[Any]]



def _trim(rows: List[List[Any]]) -> List[List[Any]]:
    out = []
    for row in rows:
        row = list(row)
        while row and (row[-1] is None or row[-1] == ""):
            row.pop()
        out.append(row)
    while out and not out[-1]:
        out.pop()
    return out


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat(sep=" ")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    return str(value)


_NUMBER = re.compile(r"^-?(0|[1-9]\d*)(\.\d+)?([eE][-+]?\d+)?$")


def _typed(value: Any) -> Any:
    """Turn numeric-looking text into numbers, keeping codes such as 007 as text."""
    if isinstance(value, str) and _NUMBER.match(value.strip()):
        text = value.strip()
        try:
            return int(text) if re.fullmatch(r"-?\d+", text) else float(text)
        except ValueError:
            return value
    return value


def _sheet_title(name: str, used: set) -> str:
    title = re.sub(r"[\[\]:*?/\\]", " ", name).strip()[:31] or "Sheet"
    base, n = title, 2
    while title.lower() in used:
        suffix = f" ({n})"
        title = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(title.lower())
    return title



def read_xlsx(path: Path) -> List[Sheet]:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), data_only=True, read_only=True)
    try:
        return [Sheet(ws.title, _trim(list(ws.iter_rows(values_only=True)))) for ws in wb.worksheets]
    finally:
        wb.close()


def read_xls(path: Path) -> List[Sheet]:
    import xlrd

    book = xlrd.open_workbook(str(path))
    sheets = []
    for sh in book.sheets():
        rows = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    row.append(bool(cell.value))
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                    row.append(None)
                else:
                    row.append(cell.value)
            rows.append(row)
        sheets.append(Sheet(sh.name, _trim(rows)))
    return sheets


def read_ods(path: Path) -> List[Sheet]:
    from odf import teletype
    from odf.opendocument import load
    from odf.table import Table, TableRow

    def repeat(node, attr) -> int:
        try:
            return max(1, int(node.getAttribute(attr) or 1))
        except (TypeError, ValueError):
            return 1

    def value(cell) -> Any:
        kind = cell.getAttribute("valuetype")
        if kind in ("float", "percentage", "currency"):
            raw = cell.getAttribute("value")
            try:
                num = float(raw)
                return int(num) if num.is_integer() else num
            except (TypeError, ValueError):
                pass
        elif kind == "date":
            raw = cell.getAttribute("datevalue") or ""
            try:
                return dt.datetime.fromisoformat(raw)
            except ValueError:
                return raw
        elif kind == "boolean":
            return cell.getAttribute("booleanvalue") == "true"
        text = "\n".join(teletype.extractText(p) for p in cell.childNodes if p.qname[1] == "p")
        return text or None

    sheets = []
    for table in load(str(path)).spreadsheet.getElementsByType(Table):
        # Rows and cells are stored run-length encoded; trailing empty runs can be
        # a million cells long, so trim runs before expanding them.
        row_runs = []
        for row in table.getElementsByType(TableRow):
            cell_runs = []
            for cell in row.childNodes:
                if cell.qname[1] in ("table-cell", "covered-table-cell"):
                    cell_runs.append((value(cell), repeat(cell, "numbercolumnsrepeated")))
            while cell_runs and cell_runs[-1][0] in (None, ""):
                cell_runs.pop()
            cells = [v for v, n in cell_runs for _ in range(n)]
            row_runs.append((cells, repeat(row, "numberrowsrepeated")))
        while row_runs and not row_runs[-1][0]:
            row_runs.pop()
        rows = [list(cells) for cells, n in row_runs for _ in range(n)]
        sheets.append(Sheet(table.getAttribute("name") or "Sheet", rows))
    return sheets


def read_delimited(path: Path, delimiter: Optional[str] = None) -> List[Sheet]:
    text = read_text(path)
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:20000], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    rows = [[_typed(v) for v in row] for row in csv.reader(text.splitlines(), delimiter=delimiter)]
    return [Sheet(path.stem, _trim(rows))]



def write_xlsx(sheets: List[Sheet], path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    for sheet in sheets or [Sheet("Sheet", [])]:
        ws = wb.create_sheet(_sheet_title(sheet.name, used))
        widths: dict = {}
        for row in sheet.rows:
            values = [_typed(v) for v in row]
            ws.append(values)
            for i, v in enumerate(values, 1):
                widths[i] = max(widths.get(i, 0), min(len(cell_text(v)), 60))
        if sheet.rows:
            for cell in ws[1]:
                cell.font = Font(bold=True)
            ws.freeze_panes = "A2"
        for i, width in widths.items():
            ws.column_dimensions[get_column_letter(i)].width = max(8, width + 2)
    wb.save(str(path))
    return path


def write_ods(sheets: List[Sheet], path: Path) -> Path:
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableCell, TableRow
    from odf.text import P

    doc = OpenDocumentSpreadsheet()
    used: set = set()
    for sheet in sheets or [Sheet("Sheet", [])]:
        table = Table(name=_sheet_title(sheet.name, used))
        for row in sheet.rows:
            tr = TableRow()
            for v in row:
                v = _typed(v)
                if v is None or v == "":
                    tc = TableCell()
                elif isinstance(v, bool):
                    tc = TableCell(valuetype="boolean", booleanvalue="true" if v else "false")
                    tc.addElement(P(text=cell_text(v)))
                elif isinstance(v, (int, float)):
                    tc = TableCell(valuetype="float", value=v)
                    tc.addElement(P(text=cell_text(v)))
                else:
                    tc = TableCell(valuetype="string")
                    tc.addElement(P(text=cell_text(v)))
                tr.addElement(tc)
            table.addElement(tr)
        doc.spreadsheet.addElement(table)
    doc.save(str(path))
    return path


def write_delimited(sheet: Sheet, path: Path, delimiter: str) -> Path:
    # utf-8 with a BOM so Excel shows accented characters correctly.
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh, delimiter=delimiter)
        for row in sheet.rows:
            writer.writerow([cell_text(v) for v in row])
    return path



def sheets_to_html(sheets: List[Sheet]) -> str:
    parts = ["<html><body>"]
    for sheet in sheets:
        if len(sheets) > 1:
            parts.append(f"<h2>{html.escape(sheet.name)}</h2>")
        if not sheet.rows:
            parts.append("<p>(empty sheet)</p>")
            continue
        width = max(len(r) for r in sheet.rows)
        parts.append("<table>")
        for i, row in enumerate(sheet.rows):
            tag = "th" if i == 0 else "td"
            cells = "".join(f"<{tag}>{html.escape(cell_text(v))}</{tag}>"
                            for v in list(row) + [None] * (width - len(row)))
            if i == 0:
                parts.append(f"<thead><tr>{cells}</tr></thead><tbody>")
            else:
                parts.append(f"<tr>{cells}</tr>")
        parts.append("</tbody></table>")
    parts.append("</body></html>")
    return "\n".join(parts)


class _TableExtractor(HTMLParser):
    """Collects tables, plus text blocks as a fallback when a document has no tables."""

    BLOCKS = {"p", "h1", "h2", "h3", "h4", "h5", "h6", "li", "pre", "blockquote", "dt", "dd"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tables: List[List[List[str]]] = []
        self.blocks: List[str] = []
        self._depth = 0
        self._row: Optional[List[str]] = None
        self._cell: Optional[List[str]] = None
        self._colspan = 1
        self._block: Optional[List[str]] = None

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self._depth += 1
            if self._depth == 1:
                self.tables.append([])
        elif self._depth == 1 and tag == "tr":
            self._row = []
        elif self._depth == 1 and tag in ("td", "th"):
            self._cell = []
            try:
                self._colspan = max(1, int(dict(attrs).get("colspan") or 1))
            except ValueError:
                self._colspan = 1
        elif tag == "br":
            (self._cell if self._cell is not None else self._block or []).append("\n")
        elif self._depth == 0 and tag in self.BLOCKS:
            self._block = []

    def handle_endtag(self, tag):
        if tag == "table":
            self._depth = max(0, self._depth - 1)
        elif self._depth == 1 and tag in ("td", "th") and self._cell is not None and self._row is not None:
            self._row.append(" ".join("".join(self._cell).split()))
            self._row.extend([""] * (self._colspan - 1))
            self._cell = None
        elif self._depth == 1 and tag == "tr" and self._row is not None:
            self.tables[-1].append(self._row)
            self._row = None
        elif self._depth == 0 and tag in self.BLOCKS and self._block is not None:
            text = " ".join("".join(self._block).split())
            if text:
                self.blocks.append(text)
            self._block = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)
        elif self._block is not None:
            self._block.append(data)


def html_to_sheets(markup: str, name: str, keep_text: bool = True) -> List[Sheet]:
    """Tables become sheets. Other text goes to a "Text" sheet when `keep_text` is set,
    or forms the only sheet when the document has no tables."""
    parser = _TableExtractor()
    parser.feed(markup)
    parser.close()
    tables = [t for t in parser.tables if t]
    text_rows = [[b] for b in parser.blocks]
    if not tables:
        return [Sheet(name, text_rows)]
    sheets = [Sheet(name if len(tables) == 1 else f"Table {i}", t) for i, t in enumerate(tables, 1)]
    if keep_text and text_rows:
        sheets.append(Sheet("Text", text_rows))
    return sheets
