"""Turn formats pandoc cannot read (PDF, PowerPoint) into HTML it can."""

from __future__ import annotations

import html
import re
from pathlib import Path
from typing import List

from .engines import ConversionError

_BULLET = re.compile("^\\s*([\u2022\u25cf\u25aa\u2013\\-*]|\\d+[.)]|[a-z][.)])\\s+")
_SENTENCE_END = re.compile(r"[.!?:;\"'”)\]]$")


def _paragraphs(page_text: str) -> List[str]:
    """Rejoin the hard-wrapped lines that PDF text extraction produces."""
    lines = [line.rstrip() for line in page_text.splitlines()]
    widths = sorted(len(line) for line in lines if line.strip())
    if not widths:
        return []
    typical = widths[int(len(widths) * 0.8)]
    paragraphs: List[str] = []
    current: List[str] = []
    for line in lines:
        stripped = line.strip()
        if not stripped:
            if current:
                paragraphs.append(" ".join(current))
                current = []
            continue
        if current and _BULLET.match(stripped):
            paragraphs.append(" ".join(current))
            current = []
        current.append(stripped)
        # A short line ending a sentence usually closes the paragraph.
        if len(stripped) < typical * 0.75 and _SENTENCE_END.search(stripped):
            paragraphs.append(" ".join(current))
            current = []
    if current:
        paragraphs.append(" ".join(current))
    return [re.sub(r"(\w)- (\w)", r"\1\2", p) for p in paragraphs]


def pdf_to_html(path: Path) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise ConversionError("This PDF is password protected.") from exc
        pages = [page.extract_text() or "" for page in reader.pages]
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError(f"Could not read this PDF: {exc}") from exc

    if not any(p.strip() for p in pages):
        raise ConversionError(
            "This PDF has no selectable text. It is probably a scanned image, "
            "which needs OCR software to convert."
        )
    body = []
    for text in pages:
        for para in _paragraphs(text):
            body.append(f"<p>{html.escape(para)}</p>")
    return "<html><body>\n" + "\n".join(body) + "\n</body></html>"


def pptx_to_html(path: Path, media_dir: Path) -> str:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    try:
        prs = Presentation(str(path))
    except Exception as exc:
        raise ConversionError(f"Could not open this PowerPoint file: {exc}") from exc

    media_dir.mkdir(parents=True, exist_ok=True)
    out = ["<html><body>"]
    image_count = 0

    def text_frame(frame) -> List[str]:
        parts: List[str] = []
        in_list = False
        for para in frame.paragraphs:
            text = "".join(run.text for run in para.runs).strip()
            if not text:
                continue
            bulleted = para.level > 0 or para._p.find(
                ".//{http://schemas.openxmlformats.org/drawingml/2006/main}buChar") is not None
            if bulleted and not in_list:
                parts.append("<ul>")
                in_list = True
            elif not bulleted and in_list:
                parts.append("</ul>")
                in_list = False
            parts.append(f"<li>{html.escape(text)}</li>" if bulleted else f"<p>{html.escape(text)}</p>")
        if in_list:
            parts.append("</ul>")
        return parts

    def shapes(items, title_shape) -> List[str]:
        nonlocal image_count
        parts: List[str] = []
        for shape in items:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                parts += shapes(shape.shapes, None)
            elif getattr(shape, "has_table", False) and shape.has_table:
                parts.append("<table>")
                for r, row in enumerate(shape.table.rows):
                    tag = "th" if r == 0 else "td"
                    parts.append("<tr>" + "".join(
                        f"<{tag}>{html.escape(cell.text)}</{tag}>" for cell in row.cells) + "</tr>")
                parts.append("</table>")
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    image = shape.image
                    image_count += 1
                    name = f"image{image_count}.{image.ext}"
                    (media_dir / name).write_bytes(image.blob)
                    parts.append(f'<p><img src="{media_dir.name}/{name}" alt=""/></p>')
                except Exception:
                    pass
            elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                parts += text_frame(shape.text_frame)
        return parts

    for number, slide in enumerate(prs.slides, 1):
        title_shape = slide.shapes.title
        heading = title_shape.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        out.append(f"<h2>{html.escape(heading or f'Slide {number}')}</h2>")
        out += shapes(slide.shapes, title_shape)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                # pandoc treats a "notes" div as speaker notes when writing slides.
                out.append(f'<div class="notes"><p>{html.escape(notes)}</p></div>')
    out.append("</body></html>")
    return "\n".join(out)
